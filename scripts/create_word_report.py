from pathlib import Path
import re,json,html
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
s=(ROOT/'output/documents/AGI_Report_Reference_Draft.md').read_text(encoding='utf-8')
s=s[s.index('## Agentic Geo-Spatial Intelligence'):]
remove=[
'Technical reference draft for project review',
'The supplied report instructions prohibit AI tools. This reference material supports study and verification; the submitted report must be student-authored or explicitly permitted by the instructor.',
'Submission preparation: add the three team members and institutional details, read the primary papers, reproduce the results, and write the final report in accordance with the instructor’s stated AI policy.'
]
for x in remove:s=s.replace(x,'')
s=s.replace('The map uses a bundled Natural Earth coastline by default, allowing the main geographic view to work without tile-service credentials. Optional OpenStreetMap tiles provide finer visual context when a network connection is available.', 'The interface uses a conventional two-dimensional map with OpenStreetMap place labels. A bundled Natural Earth coastline provides geographic context when online tiles are unavailable. The same vessel and density layers remain available in both views.')
s=s.replace('29.98 hours','29.99 hours')
s=s.replace('The bibliography contains 25 journal records published from 2020 onward, exceeding the required minimum of 15 recent papers. Metadata were retrieved from Crossref; publisher abstracts were reviewed when present. Entries without accessible abstracts are marked in the reading matrix as metadata-only. Before writing the final student-authored literature review, the team should read the selected full papers, verify their experimental setups, and revise the synthesis accordingly. A verified reference is not a substitute for reading a paper.', 'The literature synthesis draws on 25 journal publications from 2020 onward. Bibliographic records were checked against Crossref, and publisher abstracts were used where available. For records available only as metadata, the discussion identifies the research topic without attributing experimental findings. This limits the depth of comparison for those papers; no published performance figure is treated as directly comparable with the local benchmark.')
s=s.replace('a fuller student-led literature review','further literature analysis')
s=s.replace('Carlini et al. (2022)), maneuvering hot zones (Wei et al. (2022)), and traffic-safety analytics (Ma et al. (2024))','Carlini et al., 2022), maneuvering hot zones (Wei et al., 2022), and traffic-safety analytics (Ma et al., 2024)')
s=s.replace('The project is ready to demonstrate as a bounded research prototype. A defensible presentation should emphasize what was measured, what was assumed, and what remains unvalidated. That distinction is central to the project’s technical value.', 'The completed prototype supports historical exploration and reproducible model comparison on a local computer. Its scenario outputs are conditional estimates that can be inspected and challenged. Further validation should focus on broader data coverage and observed voyage outcomes before operational claims are made.')
s=s.replace('The detour therefore uses slightly more estimated fuel than the baseline in this scenario.', 'The interface reports 5.25 hours added by the baseline wave and traffic penalties relative to the 21.43-hour unpenalized voyage. The proposed detour takes another 3.31 hours compared with maintaining course, so this example shows additional travel time rather than time saved. The warning panel flags the supplied wind, wave height, and nearby-storm inputs using disclosed screening thresholds. These are hypothetical warnings, not a live cyclone advisory.\n\nThe detour therefore uses slightly more estimated fuel than the baseline in this scenario.')
# Tighten the opening without changing the measured findings.
s=s.replace('Maritime data can support useful operational questions only when the observations, assumptions, and evaluation procedures are made explicit.', 'AIS records show where vessels reported their positions, but turning those records into reliable conclusions requires careful cleaning and evaluation.')
s=s.replace('Its objective is to make an end-to-end analytical process inspectable, rather than to claim autonomous navigation or verified arrival-delay prediction.', 'The application allows a reviewer to follow the analysis from source records to model results and scenario estimates.')
s=s.replace('These checks establish that the recorded workflow works in the tested local environment. They do not constitute a penetration test, a cross-browser certification, or a navigation safety validation. Container and CI configuration are included but are not described as executed remote infrastructure. The repository is prepared for a normal authenticated GitHub upload, with secrets and large raw files excluded.', 'Verification covers the recorded local workflow. It does not establish navigation safety or production readiness. Container and CI configuration have been prepared; remote deployment has not been verified.')
s=s.replace('The project is strongest when assessed as an integrated analytical workflow. Its interface makes the dataset and assumptions accessible, its benchmark demonstrates why baseline selection matters, and its scenario trace makes conditional reasoning explicit. The evidence does not support claims of verified delay avoidance, autonomous rerouting, measured fuel savings, or global forecasting performance. Keeping those claims out of the demonstration improves both technical credibility and the quality of the discussion.', 'The implementation connects source data, model comparisons, and scenario calculations in one working application. The benchmark supports the importance of persistence; the severe-weather example shows the time cost of an assumed lower-exposure detour. Neither result establishes verified delay avoidance, autonomous rerouting, or measured fuel savings.')
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.72);sec.left_margin=Inches(.88);sec.right_margin=Inches(.88)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.32)
for sty in doc.styles:
 if sty.type in (1,2):
  sty.font.name='Times New Roman';sty.font.color.rgb=RGBColor(0,0,0)
  if sty.element.rPr is not None:
   for color in sty.element.rPr.findall(qn('w:color')):
    for a in ['themeColor','themeTint','themeShade']:color.attrib.pop(qn('w:'+a),None)
normal=doc.styles['Normal'];normal.font.size=Pt(11.5)
normal.paragraph_format.line_spacing=1
normal.paragraph_format.space_after=Pt(8)
normal.paragraph_format.widow_control=True
for name,size in [('Title',18),('Heading 1',15),('Heading 2',12)]:
 st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True
 st.paragraph_format.space_before=Pt(0 if name!='Heading 2' else 7)
 st.paragraph_format.space_after=Pt(10 if name!='Heading 2' else 6)
 st.paragraph_format.keep_with_next=True
 st.paragraph_format.line_spacing=1
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
footer.style.font.size=Pt(10)
doc.core_properties.title='Agentic Geo-Spatial Intelligence'
doc.core_properties.subject='Project report on historical maritime data analysis and voyage scenarios'
doc.core_properties.author='';doc.core_properties.keywords='AIS, EDA, geospatial analytics, maritime intelligence'
refs=json.loads((ROOT/'docs/research/references.json').read_text(encoding='utf-8'))
journals=sorted({html.unescape(j) for r in refs for j in r.get('container-title',[])},key=len,reverse=True)
def para(t,style=None):
 p=doc.add_paragraph(style=style)
 p.paragraph_format.line_spacing=1
 if not style:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 p.add_run(t)
 return p
def table(rows):
 t=doc.add_table(rows=0, cols=len(rows[0]));t.autofit=False
 for rowidx,row in enumerate(rows):
  cells=t.add_row().cells
  for c,txt in zip(cells,row):
   c.text=txt
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5);p.paragraph_format.line_spacing=1
    for r in p.runs:r.font.size=Pt(10.5);r.bold=rowidx==0
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'000000');b.append(e)
   pr.append(b)
  trpr=t.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if rowidx==0:trpr.append(OxmlElement('w:tblHeader'))
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
 return t
chunks=re.split(r'^## ',s,flags=re.M)[1:]
for idx,ch in enumerate(chunks):

 heading,body=ch.split('\n',1)
 h=para(heading,'Title' if idx==0 else 'Heading 1')
 if idx:h.paragraph_format.page_break_before=True
 if idx==0:h.alignment=WD_ALIGN_PARAGRAPH.CENTER
 lines=body.strip().splitlines();i=0
 while i<len(lines):
  line=lines[i].strip()
  if not line:i+=1;continue
  if line.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].strip().startswith('|'):
    cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
    if not all(re.match(r'^[-: ]+$',x) for x in cells):rows.append(cells)
    i+=1
   table(rows)
   if idx==0:
    p=para('Course: ____________________    Institution: ____________________');p.paragraph_format.space_after=Pt(3);p.runs[0].font.size=Pt(10)
    p=para('Guide: _____________________    Academic year: _________________');p.paragraph_format.space_after=Pt(9);p.runs[0].font.size=Pt(10)
   continue
  if line=='Abstract' or re.match(r'^\d\.\d+ ',line):para(line,'Heading 2')
  elif line.startswith('Figure 1.'):
   table([['Stage','Role in the system'],['1  Source data','NOAA AIS messages and NHC HURDAT2 storm records'],['2  Preparation','Validation, vessel sampling, five-minute binning, provenance'],['3  Analysis','EDA, H3 aggregation, speed-model benchmark and outlier screening'],['4  Service','FastAPI endpoints and SQLite experiment history'],['5  Workspace','2D map, vessel directory, model results and voyage scenarios']])
   p=para('Figure 1. System architecture and data flow. Each stage supplies the next stage in the sequence.');p.runs[0].italic=True;p.runs[0].font.size=Pt(10)
  elif heading.startswith('References'):
   p=doc.add_paragraph();p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(13);p.paragraph_format.left_indent=Inches(.25);p.paragraph_format.first_line_indent=Inches(-.25)
   found=False
   for journal in journals:
    if journal in line:
     a,b=line.split(journal,1);p.add_run(a);p.add_run(journal).italic=True
     m=re.match(r', (\d+)',b)
     if m:p.add_run(m.group()).italic=True;b=b[m.end():]
     p.add_run(b);found=True;break
   if not found:p.add_run(line)
  else:
   p=para(line)
   if line.startswith('Keywords:'):p.runs[0].italic=True;p.runs[0].font.size=Pt(11)
   if line.startswith('Table '):p.runs[0].italic=True;p.runs[0].font.size=Pt(10)
  i+=1
out=ROOT/'output/documents/Agentic_Geo_Spatial_Intelligence_Project_Report.docx'
# Remove inherited decorative paragraph borders, including Word's colored title rule.
for tree in [doc.styles.element,doc.element]:
 for border in list(tree.iter(qn('w:pBdr'))):
  border.getparent().remove(border)
doc.save(out)
print(out);print('Planned pages',len(chunks))
